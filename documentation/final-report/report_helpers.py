import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

PROJECT_DIR = r"c:\Projects\MMS"
DOC_DIR = os.path.join(PROJECT_DIR, "documentation", "final-report")
DIAGRAMS_DIR = os.path.join(DOC_DIR, "diagrams")
SCREENSHOTS_DIR = os.path.join(DOC_DIR, "screenshots")
OUTPUT_DOCX = os.path.join(PROJECT_DIR, "KJIT_Media_Management_System_Final_Project_Report.docx")
DOCX_COPY = os.path.join(DOC_DIR, "KJIT_Media_Management_System_Final_Project_Report.docx")

# Brand Colors
NAVY = RGBColor(15, 41, 66)       # #0f2942
TEXT_COLOR = RGBColor(30, 41, 59) # #1e293b
GRAY_COLOR = RGBColor(100, 116, 139) # #64748b

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def format_table(table, col_widths, headers, data, header_bg="0F2942", header_fg=RGBColor(255, 255, 255), alt_bg="F8FAFC"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], header_bg)
        set_cell_margins(hdr_cells[i], 120, 120, 140, 140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.size = Pt(9.5)
            run.font.color.rgb = header_fg
            run.font.name = "Calibri"

    # Data Rows
    for row_idx, row_data in enumerate(data):
        row_cells = table.add_row().cells
        bg = alt_bg if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].text = str(cell_value)
            set_cell_background(row_cells[col_idx], bg)
            set_cell_margins(row_cells[col_idx], 80, 80, 120, 120)
            p = row_cells[col_idx].paragraphs[0]
            if col_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.size = Pt(9)
                run.font.color.rgb = TEXT_COLOR
                run.font.name = "Calibri"

    # Set Widths
    for row in table.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = Inches(w)

def add_callout(doc, text, bold_prefix=""):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, 120, 120, 180, 180)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="0F2942"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    if bold_prefix:
        r_bold = p.add_run(bold_prefix + " ")
        r_bold.font.bold = True
        r_bold.font.size = Pt(9.5)
        r_bold.font.color.rgb = NAVY
        r_bold.font.name = "Calibri"
    r_text = p.add_run(text)
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = TEXT_COLOR
    r_text.font.name = "Calibri"
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_image_figure(doc, img_path, figure_label, caption, width=Inches(6.0)):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=width)

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(12)
        r_fig = p_cap.add_run(figure_label + ": ")
        r_fig.font.bold = True
        r_fig.font.size = Pt(9.5)
        r_fig.font.italic = True
        r_fig.font.color.rgb = NAVY
        r_fig.font.name = "Calibri"

        r_txt = p_cap.add_run(caption)
        r_txt.font.size = Pt(9.5)
        r_txt.font.italic = True
        r_txt.font.color.rgb = TEXT_COLOR
        r_txt.font.name = "Calibri"
    else:
        print(f"Warning: Image not found: {img_path}")

def add_code_snippet(doc, title, code):
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(2)
    r = p_title.add_run(title)
    r.font.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = NAVY
    r.font.name = "Calibri"

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "0F172A") # dark slate background
    set_cell_margins(cell, 120, 120, 150, 150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(code.strip())
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(226, 232, 240)
    r.font.name = "Consolas"
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

print("Report helper modules initialized.")
